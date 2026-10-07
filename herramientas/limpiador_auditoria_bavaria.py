#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Limpiador de la auditoría Bavaria (LiveTrade)
=============================================

Entrada
-------
Un Excel (o CSV) con estas columnas, en cualquier orden:

    codigo_bavaria, nombre_usuario, componente_etiqueta, componente_valor, fecha
    (opcionales: tipo_linea, actividad_id,
                 nombre_personalizado, departamento, provincia, punto_venta_id)

    componente_etiqueta = la PREGUNTA
    componente_valor    = la RESPUESTA
    tipo_linea          = a qué línea pertenece la respuesta
                          (Linea base, Inscripcion avanza, Linea de seguimiento)
    actividad_id        = actividad (formulario) de LiveTrade donde se respondió

Cada fila es una respuesta que el usuario tiene que corregir.

Qué hace
--------
1. Arregla las tildes dañadas (Â¿, Ã©, Ã‘ ...) y los espacios sobrantes.
2. Unifica la misma pregunta cuando viene escrita de varias formas
   ("5.¿Con qué...", "5. ¿Con qué...", "11. Del total..." / "4.Del total...").
3. Le pone a cada respuesta el tipo de línea al que pertenece.
4. Genera un Excel con estas hojas:
     Resumen_Usuarios     -> por usuario: cuántos códigos y respuestas debe corregir,
                             con el desglose por tipo de línea
     Usuario_x_Pregunta   -> matriz usuario + tipo de línea × pregunta (cantidad por corregir)
     Usuario_x_Codigo     -> por usuario, código bavaria y tipo de línea: qué preguntas
                             corregir y qué respuesta quedó registrada
     Resumen_Preguntas    -> por pregunta: respuestas, usuarios y códigos afectados,
                             con el desglose por tipo de línea
     Resumen_Tipo_Linea   -> por tipo de línea: respuestas, usuarios, códigos,
                             actividades y preguntas
     Preguntas_Horizontal -> una fila por código bavaria + tipo de línea,
                             una columna por pregunta
     Detalle_por_Codigo   -> toda la información limpia, ordenada por código
     Notas                -> qué se limpió y qué conviene revisar

Uso
---
    python limpiador_auditoria_bavaria.py "C:\\ruta\\auditoria_bavaria_puntos.xlsx"
    python limpiador_auditoria_bavaria.py entrada.xlsx -o salida.xlsx
    python limpiador_auditoria_bavaria.py entrada.xlsx --horizontal-por visita

    Sin argumentos busca auditoria_bavaria_puntos.xlsx en la carpeta actual.
    El resultado queda junto al original como <nombre>_limpia.xlsx.
    Si el archivo no trae la columna tipo_linea, el script igual corre y deja
    todo como "(sin tipo de línea)".

Requisitos
----------
    pip install pandas openpyxl
"""

import argparse
import datetime as dt
import numbers
import re
import sys
import time
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# CONFIGURACIÓN (lo único que normalmente hay que tocar)
# ----------------------------------------------------------------------------

# True  = la misma pregunta con distinto número se trata como una sola y se
#         muestra sin número ("11. Del total de ventas..." y "4.Del total de
#         ventas..." quedan como "Del total de ventas...").
# False = se conserva el número; solo se unen diferencias de tildes y espacios.
UNIR_PREGUNTAS_SIN_NUMERO = True

# Separador cuando hay varios valores en una misma celda.
SEPARADOR = " ; "

# Para cambiarle el nombre a una pregunta en la salida.
# Izquierda: como sale hoy en el Excel.  Derecha: como la quieres ver.
RENOMBRAR_PREGUNTAS = {
    # "¿Con qué proveedores?": "Proveedores con los que trabaja",
}

# Patrón de un usuario normal. Lo que no cumpla se lista en la hoja Notas.
PATRON_USUARIO = re.compile(r"^BAV\d+$")

# Orden en que se muestran los tipos de línea. Los que no estén aquí salen
# después, del que más respuestas tiene al que menos.
ORDEN_TIPO_LINEA = ["Linea base", "Inscripcion avanza", "Linea de seguimiento"]

OBLIGATORIAS = ["codigo_bavaria", "nombre_usuario", "componente_etiqueta", "componente_valor", "fecha"]
DATOS_LINEA = ["tipo_linea", "actividad_id"]
DATOS_PDV = ["nombre_personalizado", "departamento", "provincia", "punto_venta_id"]

SIN_PREGUNTA = "(Sin pregunta)"
SIN_CODIGO = "(sin código)"
SIN_USUARIO = "(sin usuario)"
SIN_FECHA = "(sin fecha)"
SIN_LINEA = "(sin tipo de línea)"
VACIO = "(vacío)"

# ----------------------------------------------------------------------------
# UTILIDADES DE TEXTO
# ----------------------------------------------------------------------------

_SOSPECHOSO = re.compile("[\u00c2\u00c3\u00e2]")  # Â Ã â -> señales de tildes dañadas
_NUMERACION = re.compile(r"^\s*(?:\d+(?:\.\d+)+\s+|\d+\s*[.)\-:]\s*)")  # "4." "11. " "2.1 "
_FECHA_TXT = re.compile(r"^(\d{4}-\d{2}-\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2})(?:\.\d+)?)?)?$")
_ENTERO_TXT = re.compile(r"^(0|[1-9]\d{0,14})$")  # "0123" se queda como texto


def es_vacio(v) -> bool:
    """True para None, NaN, NaT y textos en blanco."""
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def arreglar_tildes(s: str) -> str:
    """Convierte 'Â¿quÃ©' en '¿qué'. Repite por si el daño viene doble."""
    for _ in range(3):
        if not _SOSPECHOSO.search(s):
            break
        try:
            crudo = bytearray()
            for ch in s:
                try:
                    crudo += ch.encode("cp1252")
                except UnicodeEncodeError:
                    if ord(ch) < 256:  # bytes que cp1252 no define (0x81, 0x8D...)
                        crudo.append(ord(ch))
                    else:
                        raise
            nuevo = bytes(crudo).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break  # no era texto dañado: se deja como está
        if nuevo == s:
            break
        s = nuevo
    return s


def limpiar_texto(v):
    """Arregla tildes y quita espacios/tabuladores sobrantes. None si queda vacío."""
    if es_vacio(v):
        return None
    s = re.sub(r"\s+", " ", arreglar_tildes(str(v))).strip()
    return s or None


def clave(txt: str) -> str:
    """Clave para comparar preguntas: sin tildes, signos, espacios ni mayúsculas."""
    t = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", t)


def sin_numero(txt: str) -> str:
    """'11. Del total de ventas...' -> 'Del total de ventas...'"""
    resto = _NUMERACION.sub("", txt, count=1).strip()
    return resto or txt


# ----------------------------------------------------------------------------
# LIMPIEZA DE VALORES
# ----------------------------------------------------------------------------

def a_codigo(v):
    """Código bavaria / punto de venta como entero cuando se puede."""
    if es_vacio(v):
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, numbers.Integral):
        return int(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    s = limpiar_texto(v)
    return int(s) if s and _ENTERO_TXT.match(s) else s


def respuesta_limpia(v):
    """Respuesta lista para Excel: números como número, fechas como AAAA-MM-DD."""
    if es_vacio(v):
        return None
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        solo_fecha = (v.hour, v.minute, v.second) == (0, 0, 0)
        return v.strftime("%Y-%m-%d" if solo_fecha else "%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, numbers.Integral):
        return int(v)
    if isinstance(v, float):
        return int(v) if v.is_integer() else v
    s = limpiar_texto(v)
    if s is None:
        return None
    if _ENTERO_TXT.match(s):  # pasa cuando la entrada es CSV
        return int(s)
    m = _FECHA_TXT.match(s)
    if m and not any(int(x or 0) for x in m.groups()[1:]):
        return m.group(1)
    return s


def tipo_de_respuesta(v) -> str:
    """Recibe la respuesta ya limpia."""
    if v is None:
        return "Sin respuesta"
    if isinstance(v, (int, float)):
        return "Número"
    if v.lower().startswith(("http://", "https://")):
        return "Foto"
    if _FECHA_TXT.match(v):
        return "Fecha"
    if any(p.strip().lower() == "select" for p in v.split("|")):
        return "Lista con Select"
    return "Texto"


# ----------------------------------------------------------------------------
# 1) LEER
# ----------------------------------------------------------------------------

def leer_archivo(ruta: Path, hoja=0) -> pd.DataFrame:
    if ruta.suffix.lower() in (".csv", ".txt"):
        opciones = dict(dtype=str, sep=None, engine="python", keep_default_na=False)
        try:
            crudo = pd.read_csv(ruta, encoding="utf-8-sig", **opciones)
        except UnicodeDecodeError:
            crudo = pd.read_csv(ruta, encoding="cp1252", **opciones)
    else:
        crudo = pd.read_excel(ruta, sheet_name=hoja, dtype=object)

    crudo.columns = [str(c).strip().lower().lstrip("\ufeff") for c in crudo.columns]
    faltan = [c for c in OBLIGATORIAS if c not in crudo.columns]
    if faltan:
        raise ValueError(
            "Al archivo le faltan columnas: " + ", ".join(faltan)
            + "\nColumnas encontradas: " + ", ".join(map(str, crudo.columns))
        )

    crudo = crudo.astype(object)
    crudo.insert(0, "fila_excel", range(2, len(crudo) + 2))  # fila 1 = encabezado
    usadas = [c for c in OBLIGATORIAS + DATOS_LINEA + DATOS_PDV if c in crudo.columns]
    con_algo = [not all(es_vacio(v) for v in fila) for fila in crudo[usadas].itertuples(index=False, name=None)]
    return crudo.loc[con_algo, ["fila_excel"] + usadas].reset_index(drop=True)


# ----------------------------------------------------------------------------
# 2) UNIFICAR PREGUNTAS
# ----------------------------------------------------------------------------

def unificar_preguntas(etiquetas: pd.Series) -> dict:
    """Devuelve {etiqueta limpia -> nombre final de la pregunta}."""
    conteo = etiquetas.dropna().value_counts()

    def base(txt):
        return sin_numero(txt) if UNIR_PREGUNTAS_SIN_NUMERO else txt

    grupos = {}
    for txt in conteo.index:
        grupos.setdefault(clave(base(txt)), []).append(txt)

    renombres = {}
    for viejo, nuevo in RENOMBRAR_PREGUNTAS.items():
        renombres[clave(viejo)] = nuevo
        renombres[clave(sin_numero(viejo))] = nuevo

    def puntaje(txt):
        t = base(txt)
        signos_completos = t.count("¿") >= t.count("?")  # prefiere "¿qué...?" sobre "qué...?"
        return (signos_completos, conteo[txt], len(t))

    mapa = {}
    for k, variantes in grupos.items():
        nombre = base(max(variantes, key=puntaje))
        if nombre.count("(") > nombre.count(")"):
            nombre += "…"  # la etiqueta llegó cortada en el archivo
        nombre = renombres.get(k, nombre)
        for txt in variantes:
            mapa[txt] = nombre
    return mapa


def unificar_lineas(lineas: pd.Series):
    """Devuelve ({texto limpio -> nombre final del tipo de línea}, [nombres en orden]).

    Une diferencias de mayúsculas, tildes y espacios ("Línea Base" = "Linea base").
    """
    conteo = lineas.dropna().value_counts()

    grupos = {}
    for txt in conteo.index:
        grupos.setdefault(clave(txt), []).append(txt)

    mapa, total, llave_de = {}, {}, {}
    for k, variantes in grupos.items():
        nombre = max(variantes, key=lambda t: (conteo[t], t))
        total[nombre] = int(sum(conteo[t] for t in variantes))
        llave_de[nombre] = k
        for txt in variantes:
            mapa[txt] = nombre

    posicion = {clave(n): i for i, n in enumerate(ORDEN_TIPO_LINEA)}
    orden = sorted(total, key=lambda n: (posicion.get(llave_de[n], len(posicion)), -total[n], n))
    return mapa, orden


# ----------------------------------------------------------------------------
# 3) TABLA LIMPIA
# ----------------------------------------------------------------------------

def construir_detalle(crudo: pd.DataFrame):
    df = pd.DataFrame({"fila_excel": crudo["fila_excel"].astype(int)})
    celdas_arregladas = 0

    def columna_texto(nombre):
        nonlocal celdas_arregladas
        antes = list(crudo[nombre])
        despues = [limpiar_texto(v) for v in antes]
        celdas_arregladas += sum(1 for a, b in zip(antes, despues) if isinstance(a, str) and a != (b or ""))
        return pd.Series(despues, dtype=object)

    # Código y usuario: nunca vacíos, para que ninguna fila se pierda al agrupar.
    codigos = [a_codigo(v) for v in crudo["codigo_bavaria"]]
    if all(isinstance(c, int) for c in codigos):
        df["codigo_bavaria"] = codigos
    else:
        df["codigo_bavaria"] = pd.Series([SIN_CODIGO if c is None else str(c) for c in codigos], dtype=object)
    df["nombre_usuario"] = columna_texto("nombre_usuario").fillna(SIN_USUARIO)

    for col in DATOS_PDV:
        if col not in crudo.columns:
            continue
        if col == "punto_venta_id":
            df[col] = pd.Series([a_codigo(v) for v in crudo[col]], dtype=object)
        else:
            df[col] = columna_texto(col)

    df["fecha"] = pd.to_datetime(crudo["fecha"], errors="coerce").dt.normalize()

    # Tipo de línea y actividad: a qué formulario pertenece cada respuesta.
    # Nunca queda vacío, para que ninguna fila se pierda al agrupar.
    con_linea = "tipo_linea" in crudo.columns
    lineas = columna_texto("tipo_linea") if con_linea else pd.Series([None] * len(crudo), dtype=object)
    mapa_linea, orden_linea = unificar_lineas(lineas)
    df["tipo_linea"] = [mapa_linea[v] if isinstance(v, str) else SIN_LINEA for v in lineas]
    if (df["tipo_linea"] == SIN_LINEA).any() and SIN_LINEA not in orden_linea:
        orden_linea.append(SIN_LINEA)
    df["_olinea"] = df["tipo_linea"].map({n: i for i, n in enumerate(orden_linea)})

    con_actividad = "actividad_id" in crudo.columns
    if con_actividad:
        df["actividad_id"] = pd.Series([a_codigo(v) for v in crudo["actividad_id"]], dtype=object)

    # Pregunta
    df["pregunta_original"] = columna_texto("componente_etiqueta")
    mapa = unificar_preguntas(df["pregunta_original"])
    df["pregunta"] = [mapa[e] if isinstance(e, str) else SIN_PREGUNTA for e in df["pregunta_original"]]

    # Respuesta
    antes = list(crudo["componente_valor"])
    resp = [respuesta_limpia(v) for v in antes]
    celdas_arregladas += sum(1 for a, b in zip(antes, resp) if isinstance(a, str) and isinstance(b, str) and a != b)
    df["respuesta"] = pd.Series(resp, dtype=object)
    df["tipo_respuesta"] = [tipo_de_respuesta(v) for v in resp]

    # Orden de las preguntas: de la que más respuestas tiene por corregir a la que menos.
    frec = df.groupby("pregunta").size()
    orden = sorted(frec.index, key=lambda p: (p == SIN_PREGUNTA, -frec[p], p))
    df["_orden"] = df["pregunta"].map({p: i for i, p in enumerate(orden)})
    df["_fecha_txt"] = df["fecha"].dt.strftime("%Y-%m-%d").fillna(SIN_FECHA)
    df["_resp_txt"] = [VACIO if v is None else str(v) for v in resp]

    info = {
        "preguntas": orden,
        "etiquetas_originales": int(df["pregunta_original"].nunique()),
        "celdas_arregladas": int(celdas_arregladas),
        "datos_pdv": [c for c in DATOS_PDV if c in df.columns],
        "lineas": orden_linea,
        "con_linea": con_linea,
        "con_actividad": con_actividad,
    }
    return df, info


# ----------------------------------------------------------------------------
# 4) HOJAS DE SALIDA
# ----------------------------------------------------------------------------

def _unir_unicos(serie: pd.Series) -> str:
    vistos = []
    for v in serie:
        if not es_vacio(v) and v != SIN_FECHA and v not in vistos:
            vistos.append(v)
    return SEPARADOR.join(str(v) for v in vistos)


def _unir_actividades(serie: pd.Series) -> str:
    """Actividades distintas del grupo, de menor a mayor."""
    unicas = {v for v in serie if not es_vacio(v)}
    orden = sorted(unicas, key=lambda v: (isinstance(v, str), 0 if isinstance(v, str) else v, str(v)))
    return SEPARADOR.join(str(v) for v in orden)


def _conteo_por_linea(df: pd.DataFrame, llaves: list, lineas: list) -> pd.DataFrame:
    """Una columna por tipo de línea con la cantidad de respuestas (vacío = ninguna)."""
    t = df.groupby(llaves + ["tipo_linea"]).size().unstack("tipo_linea", fill_value=0)
    t = t.reindex(columns=[l for l in lineas if l in t.columns])
    return t.astype(object).where(t > 0, None)


def hoja_resumen_usuarios(df: pd.DataFrame, lineas: list) -> pd.DataFrame:
    g = df.groupby("nombre_usuario", sort=True)

    c = df.groupby(["nombre_usuario", "_orden", "pregunta"]).size().rename("n").reset_index()
    c = c.sort_values(["nombre_usuario", "n", "_orden"], ascending=[True, False, True])
    c["txt"] = c["pregunta"] + " (" + c["n"].astype(str) + ")"
    lista = c.groupby("nombre_usuario", sort=True)["txt"].agg(SEPARADOR.join)

    out = pd.concat([
        g["codigo_bavaria"].nunique().rename("codigos_bavaria_a_corregir"),
        g.size().rename("respuestas_a_corregir"),
        _conteo_por_linea(df, ["nombre_usuario"], lineas),  # respuestas por tipo de línea
        g["pregunta"].nunique().rename("preguntas_distintas"),
        g["fecha"].min().rename("primera_fecha"),
        g["fecha"].max().rename("ultima_fecha"),
        lista.rename("preguntas_a_corregir"),
    ], axis=1).reset_index()
    return out.sort_values(["respuestas_a_corregir", "nombre_usuario"], ascending=[False, True]).reset_index(drop=True)


def hoja_usuario_x_pregunta(df: pd.DataFrame, preguntas: list) -> pd.DataFrame:
    """Una fila por usuario + tipo de línea y una columna por pregunta."""
    llaves = ["nombre_usuario", "_olinea", "tipo_linea"]
    m = df.groupby(llaves + ["pregunta"]).size().unstack("pregunta", fill_value=0)
    m = m.reindex(columns=[p for p in preguntas if p in m.columns])
    total = m.sum(axis=1).astype(int)
    m = m.astype(object).where(m > 0, None)  # celda vacía = no tiene esa pregunta por corregir
    m.insert(0, "TOTAL", total)
    m = m.reset_index()

    # Primero el usuario con más respuestas por corregir; sus líneas quedan juntas.
    m["_total_usuario"] = m.groupby("nombre_usuario")["TOTAL"].transform("sum")
    m = m.sort_values(["_total_usuario", "nombre_usuario", "_olinea"], ascending=[False, True, True])
    return m.drop(columns=["_olinea", "_total_usuario"]).reset_index(drop=True)


def hoja_usuario_x_codigo(df: pd.DataFrame, datos_pdv: list, con_actividad: bool) -> pd.DataFrame:
    """Una fila por usuario + código bavaria + tipo de línea."""
    llaves = ["nombre_usuario", "codigo_bavaria", "_olinea", "tipo_linea"]
    d = df.sort_values(llaves[:3] + ["fecha", "_orden", "fila_excel"])
    g = d.groupby(llaves, sort=True)

    # "Pregunta = respuesta" (si la misma pregunta se repite, se juntan las respuestas)
    pr = (d.drop_duplicates(llaves + ["pregunta", "_resp_txt"])
          .groupby(llaves + ["_orden", "pregunta"], sort=True)["_resp_txt"].agg(" / ".join).reset_index())
    pr["txt"] = pr["pregunta"] + " = " + pr["_resp_txt"]
    lista = pr.groupby(llaves, sort=True)["txt"].agg(SEPARADOR.join)

    partes = [g[c].first() for c in datos_pdv]
    if con_actividad:
        partes.append(g["actividad_id"].agg(_unir_actividades).rename("actividades"))
    partes += [
        g["_fecha_txt"].agg(_unir_unicos).rename("fechas"),
        g.size().rename("respuestas_a_corregir"),
        g["pregunta"].nunique().rename("preguntas_distintas"),
        lista.rename("preguntas_y_respuestas_a_corregir"),
    ]
    return pd.concat(partes, axis=1).reset_index().drop(columns="_olinea")


def hoja_resumen_preguntas(df: pd.DataFrame, lineas: list) -> pd.DataFrame:
    total = len(df)
    g = df.groupby(["_orden", "pregunta"], sort=True)

    def frecuentes(grupo: pd.DataFrame) -> str:
        vc = grupo["_resp_txt"].value_counts()
        top = sorted(vc.items(), key=lambda par: (-par[1], par[0]))[:4]  # empates en orden fijo
        return SEPARADOR.join(f"{v} ({n})" for v, n in top)

    def originales(grupo: pd.DataFrame) -> str:
        """Cada forma en que viene escrita la pregunta y en qué tipo de línea aparece."""
        partes = []
        for txt, sub in grupo.dropna(subset=["pregunta_original"]).groupby("pregunta_original", sort=True):
            en = sub.sort_values("_olinea")["tipo_linea"].unique()
            partes.append(f"{txt} [{', '.join(en)}]")
        return SEPARADOR.join(partes)

    out = pd.concat([
        g.size().rename("respuestas_a_corregir"),
        _conteo_por_linea(df, ["_orden", "pregunta"], lineas),  # respuestas por tipo de línea
        g["nombre_usuario"].nunique().rename("usuarios"),
        g["codigo_bavaria"].nunique().rename("codigos_bavaria"),
    ], axis=1)
    out["pct_del_total"] = (out["respuestas_a_corregir"] / total * 100).round(1)
    out["respuestas_mas_frecuentes"] = pd.Series({k: frecuentes(v) for k, v in g}, dtype=object)
    out["como_viene_en_el_archivo"] = pd.Series({k: originales(v) for k, v in g}, dtype=object)
    return out.reset_index().drop(columns="_orden")


def hoja_resumen_lineas(df: pd.DataFrame, con_actividad: bool) -> pd.DataFrame:
    """Una fila por tipo de línea."""
    total = len(df)
    g = df.groupby(["_olinea", "tipo_linea"], sort=True)

    def actividades(grupo: pd.DataFrame) -> str:
        vc = grupo["actividad_id"].dropna().value_counts()
        pares = sorted(vc.items(), key=lambda par: str(par[0]))
        return SEPARADOR.join(f"{a} ({n})" for a, n in pares)

    def preguntas(grupo: pd.DataFrame) -> str:
        c = grupo.groupby(["_orden", "pregunta"]).size().rename("n").reset_index()
        c = c.sort_values(["n", "_orden"], ascending=[False, True])
        return SEPARADOR.join(f"{p} ({n})" for p, n in zip(c["pregunta"], c["n"]))

    out = pd.DataFrame({
        "respuestas_a_corregir": g.size(),
        "usuarios": g["nombre_usuario"].nunique(),
        "codigos_bavaria": g["codigo_bavaria"].nunique(),
        "preguntas_distintas": g["pregunta"].nunique(),
        "primera_fecha": g["fecha"].min(),
        "ultima_fecha": g["fecha"].max(),
    })
    out.insert(1, "pct_del_total", (out["respuestas_a_corregir"] / total * 100).round(1))
    if con_actividad:
        out["actividades"] = pd.Series({k: actividades(v) for k, v in g}, dtype=object)
    out["preguntas_a_corregir"] = pd.Series({k: preguntas(v) for k, v in g}, dtype=object)
    return out.reset_index().drop(columns="_olinea")


def hoja_horizontal(df: pd.DataFrame, preguntas: list, datos_pdv: list, por: str,
                    con_actividad: bool) -> pd.DataFrame:
    """Una fila por código + tipo de línea (o por visita) y una columna por pregunta."""
    llaves = ["codigo_bavaria", "_olinea", "tipo_linea"]
    if por == "visita":
        llaves += ["nombre_usuario", "_fecha_txt"]
    d = df.sort_values(["codigo_bavaria", "_olinea", "fecha", "nombre_usuario", "fila_excel"])
    g = d.groupby(llaves, sort=True)

    fijo = pd.DataFrame({c: g[c].first() for c in datos_pdv})
    if por != "visita":
        fijo["usuarios"] = g["nombre_usuario"].agg(_unir_unicos)
        fijo["fechas"] = g["_fecha_txt"].agg(_unir_unicos)
    if con_actividad:
        fijo["actividades"] = g["actividad_id"].agg(_unir_actividades)
    fijo["respuestas_a_corregir"] = g.size()

    # Si el mismo código tiene varias respuestas a la misma pregunta, van todas.
    celdas = (d.drop_duplicates(llaves + ["pregunta", "_resp_txt"])
              .groupby(llaves + ["pregunta"], sort=True)["_resp_txt"].agg(SEPARADOR.join))
    celdas = celdas.map(lambda s: s if len(s) <= 32000 else s[:32000] + " …")  # límite de celda de Excel
    ancho = celdas.unstack("pregunta")
    ancho = ancho.reindex(columns=[p for p in preguntas if p in ancho.columns])

    out = fijo.join(ancho, how="left").reset_index().drop(columns="_olinea")
    return out.rename(columns={"_fecha_txt": "fecha"})


def hoja_detalle(df: pd.DataFrame, datos_pdv: list, con_actividad: bool) -> pd.DataFrame:
    d = df.sort_values(["codigo_bavaria", "_olinea", "fecha", "nombre_usuario", "_orden", "fila_excel"])
    cols = (["codigo_bavaria", "tipo_linea"] + (["actividad_id"] if con_actividad else []) + datos_pdv
            + ["nombre_usuario", "fecha", "pregunta", "respuesta", "tipo_respuesta", "pregunta_original", "fila_excel"])
    return d[cols].reset_index(drop=True)


def hoja_notas(ruta: Path, df: pd.DataFrame, info: dict) -> pd.DataFrame:
    total = len(df)

    def pct(n):
        return f"{n:,} ({n / total * 100:.1f} %)".replace(",", ".") if total else "0"

    sin_resp = int((df["tipo_respuesta"] == "Sin respuesta").sum())
    sin_fecha = int(df["fecha"].isna().sum())
    sin_preg = int((df["pregunta"] == SIN_PREGUNTA).sum())
    raros = sorted(u for u in df["nombre_usuario"].unique() if not PATRON_USUARIO.match(str(u)))
    llave_dup = ["codigo_bavaria", "nombre_usuario", "_fecha_txt", "tipo_linea", "pregunta_original", "_resp_txt"]
    if info["con_actividad"]:
        llave_dup.append("actividad_id")
    dup = int(df.duplicated(llave_dup).sum())
    usuarios_por_cod = df.groupby("codigo_bavaria")["nombre_usuario"].nunique()
    varias_resp = int((df.groupby(["codigo_bavaria", "tipo_linea", "pregunta"])["_resp_txt"].nunique() > 1).sum())

    sin_linea = int((df["tipo_linea"] == SIN_LINEA).sum())
    con_linea = df[df["tipo_linea"] != SIN_LINEA]
    lineas_por_cod = con_linea.groupby("codigo_bavaria")["tipo_linea"].nunique()
    lineas_por_preg = con_linea.groupby("pregunta")["tipo_linea"].nunique()
    por_linea = df.groupby(["_olinea", "tipo_linea"]).size()

    fuera = []
    if "nombre_personalizado" in df.columns:
        pares = df.drop_duplicates("codigo_bavaria")[["codigo_bavaria", "nombre_personalizado"]]
        fuera = [str(c) for c, n in zip(pares["codigo_bavaria"], pares["nombre_personalizado"])
                 if str(c) not in str(n)]

    def lista_corta(items, maximo=15):
        if not items:
            return "Ninguno"
        extra = f" … y {len(items) - maximo} más" if len(items) > maximo else ""
        return ", ".join(items[:maximo]) + extra

    tipos = df["tipo_respuesta"].value_counts()
    filas = [
        ("Archivo de origen", ruta.name),
        ("Generado", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("CONTENIDO", ""),
        ("Respuestas por corregir (filas del archivo)", total),
        ("Usuarios", df["nombre_usuario"].nunique()),
        ("Códigos bavaria", df["codigo_bavaria"].nunique()),
        ("Preguntas distintas (ya unificadas)", len([p for p in info["preguntas"] if p != SIN_PREGUNTA])),
        ("Formas distintas en que venían escritas las preguntas", info["etiquetas_originales"]),
        ("Fecha mínima", df["fecha"].min().strftime("%Y-%m-%d") if df["fecha"].notna().any() else ""),
        ("Fecha máxima", df["fecha"].max().strftime("%Y-%m-%d") if df["fecha"].notna().any() else ""),
        ("", ""),
        ("TIPO DE LÍNEA", "" if info["con_linea"] else "El archivo no trae la columna tipo_linea."),
    ]
    filas += [(linea, pct(int(n))) for (_, linea), n in por_linea.items()]
    filas += [
        ("", ""),
        ("TIPO DE RESPUESTA", ""),
    ]
    filas += [(t, pct(int(n))) for t, n in tipos.items()]
    filas += [
        ("", ""),
        ("LIMPIEZA", ""),
        ("Celdas con tildes o espacios corregidos", info["celdas_arregladas"]),
        ("", ""),
        ("PARA REVISAR", ""),
        ("Filas sin respuesta", sin_resp),
        ("Filas sin fecha", sin_fecha),
        ("Filas sin pregunta", sin_preg),
        ("Filas sin tipo de línea", sin_linea),
        ("Filas repetidas idénticas (se conservan)", dup),
        ("Códigos con más de un usuario", int((usuarios_por_cod > 1).sum())),
        ("Códigos con más de un tipo de línea", int((lineas_por_cod > 1).sum())),
        ("Preguntas que aparecen en más de un tipo de línea", int((lineas_por_preg > 1).sum())),
        ("Código + tipo de línea + pregunta con más de una respuesta distinta", varias_resp),
        ("Usuarios que no tienen formato BAV###", lista_corta(raros)),
        ("Códigos que no aparecen en el nombre del punto", lista_corta(fuera)),
        ("", ""),
        ("CÓMO LEER", ""),
        ("Cada fila del archivo original", "es una respuesta por corregir."),
        ("tipo_linea", "línea a la que pertenece la respuesta; es la que dice en qué formulario se corrige."),
        ("Un código con varios tipos de línea", "sale en una fila por cada tipo de línea en Usuario_x_Codigo "
                                                "y Preguntas_Horizontal."),
        ("Columnas con nombre de tipo de línea", "en Resumen_Usuarios y Resumen_Preguntas: cuántas respuestas "
                                                 "por corregir hay en esa línea (vacío = ninguna)."),
        (SIN_LINEA, "la fila no traía tipo de línea en el archivo original."),
        ("Preguntas", "van sin número porque la misma pregunta cambia de número según el formulario; "
                      "el texto original está en Detalle_por_Codigo (pregunta_original) y, con su tipo de "
                      "línea, en Resumen_Preguntas (como_viene_en_el_archivo)."
         if UNIR_PREGUNTAS_SIN_NUMERO else "conservan el número con el que vienen en el archivo."),
        ("Pregunta terminada en '…'", "llegó cortada en el archivo original."),
        (VACIO, "la fila no traía respuesta."),
        ("Varias respuestas en una celda", f"separadas por '{SEPARADOR.strip()}': el código tiene más de una "
                                           "respuesta a esa pregunta en la misma línea (otra fecha u otro usuario)."),
    ]
    return pd.DataFrame(filas, columns=["Concepto", "Valor"])


# ----------------------------------------------------------------------------
# 5) ESCRIBIR EL EXCEL
# ----------------------------------------------------------------------------

_FUENTE_ENC = Font(bold=True, color="FFFFFF")
_FONDO_ENC = PatternFill("solid", fgColor="1F3864")
_ALIN_ENC = Alignment(vertical="center", wrap_text=True)

_ANCHOS = {  # anchos fijos por nombre de columna; el resto se calcula
    "preguntas_a_corregir": 100, "preguntas_y_respuestas_a_corregir": 100,
    "respuestas_mas_frecuentes": 60, "como_viene_en_el_archivo": 70,
    "nombre_personalizado": 40, "pregunta": 60, "pregunta_original": 60, "respuesta": 34,
    "fechas": 24, "usuarios": 18, "Concepto": 62, "Valor": 90,
    "tipo_linea": 22,
}


def _celda(v):
    """Convierte cada valor a algo que openpyxl pueda escribir."""
    if v is None:
        return None
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.date()
    if isinstance(v, float) and v != v:
        return None
    if hasattr(v, "item"):  # tipos de numpy
        return v.item()
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v


def escribir_hoja(wb: Workbook, nombre: str, df: pd.DataFrame, congelar: str = "A2",
                  alto_enc: float = 30, cols_pregunta=(), ancho_pregunta: float = 26, filtro: bool = True):
    ws = wb.create_sheet(nombre)
    ws.freeze_panes = congelar
    ws.row_dimensions[1].height = alto_enc

    muestra = df.head(300)
    for i, col in enumerate(df.columns, start=1):
        if col in _ANCHOS:
            ancho = _ANCHOS[col]
        elif col in cols_pregunta:
            ancho = ancho_pregunta
        else:
            largo = max([len(str(col))] + [len(str(v)) for v in muestra[col] if not es_vacio(v)])
            ancho = min(max(largo + 2, 10), 42)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    if filtro and len(df.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"

    enc = []
    for col in df.columns:
        c = WriteOnlyCell(ws, value=str(col))
        c.font, c.fill, c.alignment = _FUENTE_ENC, _FONDO_ENC, _ALIN_ENC
        enc.append(c)
    ws.append(enc)

    for fila in df.astype(object).itertuples(index=False, name=None):
        ws.append([_celda(v) for v in fila])


def escribir_excel(salida: Path, hojas: list):
    wb = Workbook(write_only=True)
    for nombre, df, opciones in hojas:
        escribir_hoja(wb, nombre, df, **opciones)
    wb.save(salida)


# ----------------------------------------------------------------------------
# PRINCIPAL
# ----------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Limpiador de la auditoría Bavaria (LiveTrade).")
    ap.add_argument("entrada", nargs="?", default="auditoria_bavaria_puntos.xlsx",
                    help="Excel o CSV de la auditoría (por defecto: auditoria_bavaria_puntos.xlsx)")
    ap.add_argument("-o", "--salida", help="Excel de salida (por defecto: <entrada>_limpia.xlsx)")
    ap.add_argument("--hoja", default=0, help="Nombre o posición de la hoja a leer (por defecto la primera)")
    ap.add_argument("--horizontal-por", choices=["codigo", "visita"], default="codigo",
                    help="'codigo' = una fila por código bavaria + tipo de línea (por defecto); "
                         "'visita' = una fila por código + tipo de línea + usuario + fecha")
    args = ap.parse_args(argv)

    try:  # evita errores de consola en Windows con caracteres especiales
        sys.stdout.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

    entrada = Path(args.entrada)
    if not entrada.exists():
        print(f"No encuentro el archivo: {entrada}", file=sys.stderr)
        return 1
    salida = Path(args.salida) if args.salida else entrada.with_name(entrada.stem + "_limpia.xlsx")
    hoja = int(args.hoja) if str(args.hoja).isdigit() else args.hoja

    t0 = time.time()
    print(f"[1/4] Leyendo {entrada.name} ...")
    try:
        crudo = leer_archivo(entrada, hoja)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        return 1
    if crudo.empty:
        print("El archivo no tiene filas con datos.", file=sys.stderr)
        return 1

    print(f"[2/4] Limpiando {len(crudo):,} filas ...".replace(",", "."))
    df, info = construir_detalle(crudo)
    preguntas, datos_pdv = info["preguntas"], info["datos_pdv"]
    lineas, con_actividad = info["lineas"], info["con_actividad"]
    if not info["con_linea"]:
        print(f"      Aviso: el archivo no trae la columna tipo_linea; todo queda como {SIN_LINEA}.")

    print("[3/4] Armando hojas ...")
    resumen_usr = hoja_resumen_usuarios(df, lineas)
    usr_preg = hoja_usuario_x_pregunta(df, preguntas)
    usr_cod = hoja_usuario_x_codigo(df, datos_pdv, con_actividad)
    resumen_preg = hoja_resumen_preguntas(df, lineas)
    resumen_lin = hoja_resumen_lineas(df, con_actividad)
    horizontal = hoja_horizontal(df, preguntas, datos_pdv, args.horizontal_por, con_actividad)
    detalle = hoja_detalle(df, datos_pdv, con_actividad)
    notas = hoja_notas(entrada, df, info)

    # Columnas que quedan fijas al desplazarse: las llaves y el nombre del punto.
    fijas_h = (4 if args.horizontal_por == "visita" else 2) + ("nombre_personalizado" in datos_pdv)
    hojas = [
        ("Resumen_Usuarios", resumen_usr, {"congelar": "B2"}),
        ("Usuario_x_Pregunta", usr_preg, {"congelar": "D2", "alto_enc": 95, "cols_pregunta": preguntas,
                                          "ancho_pregunta": 17}),
        ("Usuario_x_Codigo", usr_cod, {"congelar": "D2"}),
        ("Resumen_Preguntas", resumen_preg, {"congelar": "B2"}),
        ("Resumen_Tipo_Linea", resumen_lin, {"congelar": "B2"}),
        ("Preguntas_Horizontal", horizontal, {"congelar": f"{get_column_letter(fijas_h + 1)}2", "alto_enc": 80,
                                              "cols_pregunta": preguntas}),
        ("Detalle_por_Codigo", detalle, {"congelar": "C2"}),
        ("Notas", notas, {"filtro": False}),
    ]

    print(f"[4/4] Escribiendo {salida.name} ...")
    try:
        escribir_excel(salida, hojas)
    except PermissionError:
        print(f"No pude guardar {salida}. Si lo tienes abierto en Excel, ciérralo y vuelve a correr.",
              file=sys.stderr)
        return 1

    n = lambda x: f"{x:,}".replace(",", ".")
    print()
    print(f"Listo en {time.time() - t0:.0f} s  ->  {salida}")
    print(f"  Respuestas por corregir: {n(len(df))}")
    print(f"  Usuarios:                {n(df['nombre_usuario'].nunique())}")
    print(f"  Códigos bavaria:         {n(df['codigo_bavaria'].nunique())}")
    print(f"  Preguntas distintas:     {n(len(preguntas))}")
    print()
    print("Respuestas por corregir según tipo de línea:")
    print(resumen_lin[["tipo_linea", "respuestas_a_corregir", "usuarios", "codigos_bavaria"]].to_string(index=False))
    print()
    print("Usuarios con más respuestas por corregir:")
    print(resumen_usr.head(10)[["nombre_usuario", "codigos_bavaria_a_corregir", "respuestas_a_corregir"]]
          .to_string(index=False))
    print()
    print("Preguntas con más respuestas por corregir:")
    print(resumen_preg.head(8)[["pregunta", "respuestas_a_corregir", "usuarios"]].to_string(index=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())