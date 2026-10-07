"""Compara dos Excel celda por celda y en formato. Uso: comparar.py <referencia.xlsx> <port.xlsx>"""
import datetime as dt, sys
from openpyxl import load_workbook
a, b = load_workbook(sys.argv[1]), load_workbook(sys.argv[2])
dif = []
def norm(v):
    if isinstance(v, dt.datetime): return ("fecha", v.date().isoformat(), )
    if isinstance(v, float) and v.is_integer(): return int(v)
    if v == "": return None
    return v
if a.sheetnames != b.sheetnames: dif.append(("libro", "hojas", a.sheetnames, b.sheetnames))
celdas = 0
for nombre in a.sheetnames:
    x, y = a[nombre], b[nombre]
    if (x.max_row, x.max_column) != (y.max_row, y.max_column):
        dif.append((nombre, "tamaño", (x.max_row, x.max_column), (y.max_row, y.max_column)))
    for i, (fx, fy) in enumerate(zip(x.iter_rows(), y.iter_rows()), start=1):
        for cx, cy in zip(fx, fy):
            celdas += 1
            if nombre == "Notas" and fx[0].value == "Generado" and cx.column == 2:  # hora de cada corrida
                continue
            if norm(cx.value) != norm(cy.value):
                dif.append((nombre, cx.coordinate, repr(cx.value)[:80], repr(cy.value)[:80]))
            elif isinstance(cx.value, dt.datetime) and cx.number_format != cy.number_format:
                dif.append((nombre, cx.coordinate, "formato " + cx.number_format, cy.number_format))
    if x.freeze_panes != y.freeze_panes: dif.append((nombre, "congelar", x.freeze_panes, y.freeze_panes))
    if x.auto_filter.ref != y.auto_filter.ref: dif.append((nombre, "filtro", x.auto_filter.ref, y.auto_filter.ref))
    if x.row_dimensions[1].height != y.row_dimensions[1].height:
        dif.append((nombre, "alto fila 1", x.row_dimensions[1].height, y.row_dimensions[1].height))
    def anchos(ws):  # <col min max> agrupa columnas seguidas con el mismo ancho
        r = {}
        for d in ws.column_dimensions.values():
            for k in range(d.min, d.max + 1): r[k] = d.width
        return r
    ax, ay = anchos(x), anchos(y)
    for k in sorted(set(ax) | set(ay)):
        if ax.get(k) != ay.get(k): dif.append((nombre, f"ancho col {k}", ax.get(k), ay.get(k)))
    for cx, cy in zip(x[1], y[1]):
        ex = (cx.font.b, (cx.font.color.rgb or "")[-6:], (cx.fill.fgColor.rgb or "")[-6:], cx.fill.fill_type, cx.alignment.wrap_text, cx.alignment.vertical)
        ey = (cy.font.b, (cy.font.color.rgb or "")[-6:], (cy.fill.fgColor.rgb or "")[-6:], cy.fill.fill_type, cy.alignment.wrap_text, cy.alignment.vertical)
        if ex != ey: dif.append((nombre, f"estilo {cx.coordinate}", ex, ey)); break
print(f"{celdas:,} celdas en {len(a.sheetnames)} hojas | diferencias: {len(dif)}")
for d in dif[:25]: print("  ", d)
